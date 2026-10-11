"""
Tests de los formatos de salida distintos del PDF.

Lo que se comprueba aquí no es que los ficheros existan, sino lo que hace
falta para poder defenderlos: que las cuatro copias dicen lo mismo, llevan el
mismo código de informe y conservan la trazabilidad evento → registro de
origen. Un DOCX que se hubiera compuesto en otro momento y dijera algo
distinto que el PDF firmado sería peor que no tenerlo.
"""
import io
import json
from unittest.mock import patch

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.modules.reports import repository
from app.modules.reports.docx_builder import construir_docx
from app.modules.reports.html_builder import construir_html
from app.modules.reports.json_builder import construir_json
from app.modules.reports.pdf_builder import construir_pdf
from tests.conftest import NARRATIVA_FALSA, documento_de_prueba

_SECCIONES = (
    "Objeto y alcance",
    "Antecedentes",
    "Material analizado",
    "Metodología",
    "Cronología de eventos",
    "Hallazgos",
    "Conclusiones",
    "inteligencia artificial",
    "Firma",
    "Anexo A",
    "Anexo B",
    "Anexo C",
)


def _texto_pdf(contenido: bytes) -> str:
    lector = PdfReader(io.BytesIO(contenido))
    return "\n".join(pagina.extract_text() or "" for pagina in lector.pages)


def _texto_docx(contenido: bytes) -> str:
    docx = DocxDocument(io.BytesIO(contenido))
    partes = [parrafo.text for parrafo in docx.paragraphs]
    for tabla in docx.tables:
        partes += [celda.text for fila in tabla.rows for celda in fila.cells]
    return "\n".join(partes)


# --- DOCX -------------------------------------------------------------------


def test_docx_lleva_las_secciones_periciales(db_session, case_con_anexos):
    texto = _texto_docx(construir_docx(documento_de_prueba(db_session, case_con_anexos)))

    for seccion in _SECCIONES:
        assert seccion in texto, f"falta la sección «{seccion}»"


def test_docx_lleva_la_cronologia_con_su_referencia(db_session, case_con_timeline):
    texto = _texto_docx(construir_docx(documento_de_prueba(db_session, case_con_timeline)))

    assert "tracking#0" in texto
    assert "Ref." in texto


def test_docx_numera_las_paginas_con_campos_de_word(db_session, case_con_timeline):
    """
    Un número de página fijo mentiría en cuanto alguien repaginase el
    documento, que es justo lo que se hace al editar un DOCX.
    """
    docx = DocxDocument(
        io.BytesIO(construir_docx(documento_de_prueba(db_session, case_con_timeline)))
    )
    pie = docx.sections[0].footer.paragraphs[0]
    xml = pie._p.xml

    assert "PAGE" in xml and "NUMPAGES" in xml
    assert f"INF-{case_con_timeline.id:04d}-v1" in pie.text


def test_caracteres_xml_no_rompen_el_docx(db_session, case_con_timeline):
    eventos = repository.get_timeline_events(db_session, case_con_timeline.id)
    eventos[0].descripcion = 'Payload <script>alert("x")</script> & "comillas"'

    texto = _texto_docx(
        construir_docx(documento_de_prueba(db_session, case_con_timeline, eventos))
    )

    assert "<script>" in texto


# --- HTML -------------------------------------------------------------------


def test_html_lleva_las_secciones_periciales(db_session, case_con_anexos):
    html = construir_html(documento_de_prueba(db_session, case_con_anexos)).decode()

    for seccion in _SECCIONES:
        assert seccion in html, f"falta la sección «{seccion}»"


def test_html_es_autonomo(db_session, case_con_timeline):
    """Sin recursos externos: tiene que seguir leyéndose en una máquina sin red."""
    html = construir_html(documento_de_prueba(db_session, case_con_timeline)).decode()

    assert "<style>" in html
    assert "http://" not in html and "https://" not in html
    assert "<link" not in html and "<script" not in html


def test_html_escapa_el_texto_de_la_base_de_datos(db_session, case_con_timeline):
    """Sin escapar, una descripción con etiquetas se ejecutaría al abrir el informe."""
    eventos = repository.get_timeline_events(db_session, case_con_timeline.id)
    eventos[0].descripcion = '<script>alert("x")</script> & <b>negrita</b>'

    html = construir_html(
        documento_de_prueba(db_session, case_con_timeline, eventos)
    ).decode()

    assert "&lt;script&gt;" in html
    assert "<script>alert" not in html


# --- JSON -------------------------------------------------------------------


def test_json_declara_los_metadatos_del_informe(db_session, case_con_anexos):
    contenido = json.loads(construir_json(documento_de_prueba(db_session, case_con_anexos)))

    assert contenido["formato_version"] == "1"
    informe = contenido["informe"]
    assert informe["codigo"] == f"INF-{case_con_anexos.id:04d}-v1"
    assert informe["case_id"] == case_con_anexos.id
    assert informe["version"] == 1
    assert informe["generado_por"] == "Ana Pérez"
    assert informe["modelo_ia"] == NARRATIVA_FALSA.modelo


def test_json_conserva_la_cronologia_y_su_trazabilidad(db_session, case_con_timeline):
    contenido = json.loads(
        construir_json(documento_de_prueba(db_session, case_con_timeline))
    )
    eventos = repository.get_timeline_events(db_session, case_con_timeline.id)

    cronologia = next(
        seccion
        for seccion in contenido["secciones"]
        if "Cronología" in seccion["titulo"]
    )
    tabla = cronologia["tablas"][0]

    assert tabla["cabeceras"][0] == "Ref."
    assert len(tabla["filas"]) == len(eventos)
    assert tabla["filas"][0][0] == "E1"
    assert tabla["filas"][0][2] == "tracking#0"


def test_json_recoge_todas_las_secciones(db_session, case_con_anexos):
    contenido = json.loads(construir_json(documento_de_prueba(db_session, case_con_anexos)))
    titulos = " | ".join(seccion["titulo"] for seccion in contenido["secciones"])

    for seccion in _SECCIONES:
        assert seccion in titulos, f"falta la sección «{seccion}»"


# --- Coherencia entre copias ------------------------------------------------


def test_las_cuatro_copias_declaran_el_mismo_codigo(db_session, case_con_anexos):
    """
    Es la comprobación que sostiene todo lo demás: cuatro ficheros que se
    presentan como el mismo informe tienen que poder contrastarse entre sí.
    """
    doc = documento_de_prueba(db_session, case_con_anexos)
    codigo = f"INF-{case_con_anexos.id:04d}-v1"

    assert codigo in _texto_pdf(construir_pdf(doc))
    assert codigo in _texto_docx(construir_docx(doc))
    assert codigo in construir_html(doc).decode()
    assert codigo in json.loads(construir_json(doc))["informe"]["codigo"]


# --- Emisión y descarga -----------------------------------------------------


def test_generate_report_emite_los_cuatro_formatos(client, case_con_timeline, storage_dir):
    with patch("app.modules.reports.ai_narrative.generar_narrativa", return_value=NARRATIVA_FALSA):
        respuesta = client.post(
            "/reports/generate",
            json={"case_id": case_con_timeline.id, "generado_por": "Tester"},
        )
    assert respuesta.status_code == 200

    case_dir = storage_dir / f"case_{case_con_timeline.id}"
    for extension in ("pdf", "docx", "html", "json"):
        assert (case_dir / f"report_v1.{extension}").is_file(), f"falta el .{extension}"


def test_download_sirve_cada_formato_con_su_media_type(client, case_con_timeline, storage_dir):
    with patch("app.modules.reports.ai_narrative.generar_narrativa", return_value=NARRATIVA_FALSA):
        report = client.post(
            "/reports/generate",
            json={"case_id": case_con_timeline.id, "generado_por": "Tester"},
        ).json()

    esperado = {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "html": "text/html; charset=utf-8",
        "json": "application/json",
    }
    for formato, media_type in esperado.items():
        respuesta = client.get(f"/reports/{report['id']}/download?formato={formato}")
        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"] == media_type
        assert formato in respuesta.headers["content-disposition"]


def test_el_html_se_sirve_en_linea_y_el_resto_como_descarga(client, case_con_timeline, storage_dir):
    """
    El HTML existe para leerlo en el navegador: como `attachment` se descarga y
    hay que abrirlo a mano, que es justo lo que ese formato venía a evitar.
    Los otros tres se descargan; nadie quiere un PDF pintado dentro de la
    pestaña de una API.
    """
    with patch("app.modules.reports.ai_narrative.generar_narrativa", return_value=NARRATIVA_FALSA):
        report = client.post(
            "/reports/generate",
            json={"case_id": case_con_timeline.id, "generado_por": "Tester"},
        ).json()

    html = client.get(f"/reports/{report['id']}/download?formato=html")
    assert html.headers["content-disposition"].startswith("inline")

    for formato in ("pdf", "docx", "json"):
        respuesta = client.get(f"/reports/{report['id']}/download?formato={formato}")
        assert respuesta.headers["content-disposition"].startswith("attachment")


def test_download_sin_formato_sigue_dando_el_pdf(client, case_con_timeline, storage_dir):
    """La firma anterior del endpoint no se rompe: quien no pide formato, recibe PDF."""
    with patch("app.modules.reports.ai_narrative.generar_narrativa", return_value=NARRATIVA_FALSA):
        report = client.post(
            "/reports/generate",
            json={"case_id": case_con_timeline.id, "generado_por": "Tester"},
        ).json()

    respuesta = client.get(f"/reports/{report['id']}/download")

    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/pdf"


def test_download_de_un_formato_inexistente_da_422(client, case_con_timeline, storage_dir):
    with patch("app.modules.reports.ai_narrative.generar_narrativa", return_value=NARRATIVA_FALSA):
        report = client.post(
            "/reports/generate",
            json={"case_id": case_con_timeline.id, "generado_por": "Tester"},
        ).json()

    respuesta = client.get(f"/reports/{report['id']}/download?formato=odt")

    assert respuesta.status_code == 422


def test_download_de_una_copia_borrada_da_410(client, case_con_timeline, storage_dir):
    """
    Un informe emitido antes de que el módulo generase varios formatos solo
    tiene PDF. No se compone uno nuevo al vuelo: no coincidiría con el
    documento firmado en su día.
    """
    with patch("app.modules.reports.ai_narrative.generar_narrativa", return_value=NARRATIVA_FALSA):
        report = client.post(
            "/reports/generate",
            json={"case_id": case_con_timeline.id, "generado_por": "Tester"},
        ).json()

    (storage_dir / f"case_{case_con_timeline.id}" / "report_v1.docx").unlink()

    respuesta = client.get(f"/reports/{report['id']}/download?formato=docx")

    assert respuesta.status_code == 410
    assert respuesta.json()["detail"].startswith(f"El informe {report['id']}")
