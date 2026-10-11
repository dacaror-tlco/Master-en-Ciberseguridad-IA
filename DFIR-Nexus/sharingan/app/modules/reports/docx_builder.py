"""
Renderizado del informe a DOCX (python-docx).

Es el formato que suele pedir quien recibe el informe —abogado, instructor,
otro perito— porque necesita citarlo, anotarlo o incorporarlo a un escrito
mayor. El PDF sigue siendo el documento que se firma; el DOCX es la copia
manipulable, y por eso lleva el mismo código de informe y la misma fecha de
emisión: cualquiera puede comprobar de qué original procede.

El pie usa campos de Word (`PAGE`, `NUMPAGES`) en lugar de números fijos.
Word los recalcula al repaginar, que es lo que hace un editor cuando alguien
cambia márgenes o tipografía; un número escrito a mano mentiría en cuanto se
tocara el documento.
"""
import io

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.modules.reports.document import (
    Documento,
    Espacio,
    Parrafo,
    SaltoPagina,
    Tabla,
    Titulo,
)

_CABECERA_FONDO = "20242F"


def construir_docx(documento: Documento) -> bytes:
    docx = DocxDocument()
    _metadatos(docx, documento)
    _pie(docx, documento)

    for bloque in documento.bloques:
        match bloque:
            case Titulo(texto=texto, nivel=nivel):
                parrafo = docx.add_heading(texto, level=nivel)
                if nivel == 0:
                    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            case Parrafo(texto=texto):
                parrafo = docx.add_paragraph(str(texto))
                parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            case Tabla():
                _tabla(docx, bloque)
            case Espacio(alto_cm=alto):
                docx.add_paragraph().paragraph_format.space_after = Cm(alto)
            case SaltoPagina():
                docx.add_page_break()

    buffer = io.BytesIO()
    docx.save(buffer)
    return buffer.getvalue()


def _metadatos(docx, documento: Documento) -> None:
    propiedades = docx.core_properties
    propiedades.title = documento.titulo
    propiedades.author = documento.generado_por
    propiedades.subject = f"Informe pericial del caso {documento.case_id}"
    propiedades.comments = (
        f"Generado por DFIR Nexus — Módulo 5 (Sharingan). Código {documento.codigo}."
    )


def _tabla(docx, bloque: Tabla) -> None:
    tabla = docx.add_table(rows=1, cols=len(bloque.cabeceras))
    tabla.style = "Table Grid"

    for celda, texto in zip(tabla.rows[0].cells, bloque.cabeceras):
        _sombrear(celda, _CABECERA_FONDO)
        parrafo = celda.paragraphs[0]
        tramo = parrafo.add_run(texto)
        tramo.bold = True
        tramo.font.size = Pt(9)
        tramo.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for fila in bloque.filas:
        celdas = tabla.add_row().cells
        for celda, texto in zip(celdas, fila):
            parrafo = celda.paragraphs[0]
            parrafo.add_run(str(texto)).font.size = Pt(9)

    docx.add_paragraph()


def _sombrear(celda, color_hex: str) -> None:
    """Fondo de celda. python-docx no lo expone, hay que bajar al XML."""
    sombreado = OxmlElement("w:shd")
    sombreado.set(qn("w:val"), "clear")
    sombreado.set(qn("w:fill"), color_hex)
    celda._tc.get_or_add_tcPr().append(sombreado)


def _pie(docx, documento: Documento) -> None:
    """Código de informe a la izquierda y «Página X de Y» a la derecha."""
    parrafo = docx.sections[0].footer.paragraphs[0]
    parrafo.text = ""

    margenes = docx.sections[0]
    ancho_util = margenes.page_width - margenes.left_margin - margenes.right_margin
    parrafo.paragraph_format.tab_stops.add_tab_stop(ancho_util, WD_TAB_ALIGNMENT.RIGHT)

    parrafo.add_run(f"{documento.codigo}\tPágina ")
    _campo(parrafo, "PAGE")
    parrafo.add_run(" de ")
    _campo(parrafo, "NUMPAGES")

    for tramo in parrafo.runs:
        tramo.font.size = Pt(8)
        tramo.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)


def _campo(parrafo, instruccion: str) -> None:
    """Inserta un campo de Word (se recalcula solo al repaginar)."""
    tramo = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    texto = OxmlElement("w:instrText")
    texto.set(qn("xml:space"), "preserve")
    texto.text = f" {instruccion} "
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for nodo in (inicio, texto, fin):
        tramo._r.append(nodo)
